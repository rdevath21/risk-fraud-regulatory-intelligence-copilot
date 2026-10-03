from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Colors
BG = RGBColor(15, 23, 42)
WHITE = RGBColor(226, 232, 240)
GRAY = RGBColor(148, 163, 184)
DARK_GRAY = RGBColor(71, 85, 105)
BLUE = RGBColor(56, 189, 248)
PURPLE = RGBColor(168, 85, 247)
GREEN = RGBColor(52, 211, 153)
ORANGE = RGBColor(251, 146, 60)
RED = RGBColor(248, 113, 113)
CARD_BG = RGBColor(20, 30, 55)
CARD_BORDER = RGBColor(40, 50, 80)

def set_bg(slide):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = BG

def add_text_box(slide, left, top, width, height, text, font_size=14, color=WHITE, bold=False, alignment=PP_ALIGN.LEFT, font_name='Calibri'):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return txBox

def add_card(slide, left, top, width, height, fill_color=CARD_BG):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.color.rgb = CARD_BORDER
    shape.line.width = Pt(1)
    shape.shadow.inherit = False
    return shape

def add_colored_card(slide, left, top, width, height, border_color):
    shape = add_card(slide, left, top, width, height)
    shape.line.color.rgb = border_color
    shape.line.width = Pt(2)
    return shape

# ═══════════════════════════════════════════════
# SLIDE 1: TITLE
# ═══════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_text_box(slide, 1.5, 0.5, 10, 0.5, "CoCo Hackathon GCC Edition 2026", 14, BLUE, True)
add_text_box(slide, 1.5, 1.3, 10, 1.5, "Risk, Fraud &\nRegulatory Intelligence Copilot", 44, WHITE, True)
add_text_box(slide, 1.5, 3.3, 10, 1, "AI-powered compliance investigation platform that surfaces risk signals,\ngenerates audit-ready evidence, and produces regulatory outputs\nfrom natural language questions.", 18, GRAY)
add_text_box(slide, 1.5, 4.8, 10, 0.5, "Snowflake Cortex AI   |   CoCo Desktop   |   Streamlit   |   MCP Server   |   Cortex Agents", 14, DARK_GRAY)
add_text_box(slide, 1.5, 5.6, 10, 0.4, "Built end-to-end with CoCo Desktop  |  40+ Snowflake Objects  |  3 Cortex Agents  |  5 Surfaces", 13, DARK_GRAY)
add_text_box(slide, 1.5, 6.2, 10, 0.4, "GitHub: github.com/rdevath21/risk-fraud-regulatory-intelligence-copilot", 12, DARK_GRAY)
add_text_box(slide, 12, 0.3, 1, 0.3, "01 / 08", 10, DARK_GRAY, alignment=PP_ALIGN.RIGHT)

# ═══════════════════════════════════════════════
# SLIDE 2: PROBLEM BRIEF
# ═══════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_text_box(slide, 0.8, 0.3, 3, 0.3, "PROBLEM BRIEF", 11, RED, True)
add_text_box(slide, 12, 0.3, 1, 0.3, "02 / 08", 10, DARK_GRAY, alignment=PP_ALIGN.RIGHT)
add_text_box(slide, 0.8, 0.7, 10, 0.6, "The Problem We Solve", 32, WHITE, True)

# Left column - Pain points
add_colored_card(slide, 0.8, 1.5, 5.8, 2.2, RED)
add_text_box(slide, 1.1, 1.6, 5.3, 0.4, "Industry Context", 16, RED, True)
add_text_box(slide, 1.1, 2.1, 5.3, 1.4, "Banking and NBFC teams manage real-time fraud detection, liquidity risk, credit risk, and regulatory reporting (AML, Basel, local regulations). These processes are largely manual today, creating compliance gaps and investigation backlogs.", 13, GRAY)

add_colored_card(slide, 0.8, 3.9, 5.8, 3, RED)
add_text_box(slide, 1.1, 4.0, 5.3, 0.3, "Current Pain Points", 16, RED, True)
add_text_box(slide, 1.1, 4.5, 5.3, 2.2,
    "X  4-8 hours per alert for manual investigation and evidence compilation\n\n"
    "X  Regulatory policies scattered across dozens of documents with no unified search\n\n"
    "X  No audit trail for AI-assisted decisions - regulators demand explainability\n\n"
    "X  SAR filing backlogs: 30-day deadline often at risk due to slow evidence generation",
    13, GRAY)

# Right column - Personas
add_text_box(slide, 7.2, 1.5, 5, 0.3, "TARGET USERS", 11, GRAY, True)

add_colored_card(slide, 7.2, 2.0, 5.5, 1.3, BLUE)
add_text_box(slide, 7.5, 2.1, 5, 0.3, "AML Compliance Analyst", 15, WHITE, True)
add_text_box(slide, 7.5, 2.5, 5, 0.7, "Triages alerts daily, investigates suspicious transactions, generates SAR evidence packages. Needs fast access to transaction data + regulatory context.", 12, GRAY)

add_colored_card(slide, 7.2, 3.5, 5.5, 1.3, PURPLE)
add_text_box(slide, 7.5, 3.6, 5, 0.3, "BSA Officer / Compliance Manager", 15, WHITE, True)
add_text_box(slide, 7.5, 4.0, 5, 0.7, "Reviews investigations, makes SAR filing decisions, reports to board. Needs governed, auditable outputs with regulatory citations.", 12, GRAY)

add_colored_card(slide, 7.2, 5.0, 5.5, 1.3, GREEN)
add_text_box(slide, 7.5, 5.1, 5, 0.3, "Risk Auditor", 15, WHITE, True)
add_text_box(slide, 7.5, 5.5, 5, 0.7, "Performs periodic reviews of investigation quality and AI usage. Needs read-only access with full audit trail visibility.", 12, GRAY)

# ═══════════════════════════════════════════════
# SLIDE 3: ARCHITECTURE
# ═══════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_text_box(slide, 0.8, 0.3, 3, 0.3, "ARCHITECTURE DIAGRAM", 11, BLUE, True)
add_text_box(slide, 12, 0.3, 1, 0.3, "03 / 08", 10, DARK_GRAY, alignment=PP_ALIGN.RIGHT)
add_text_box(slide, 0.8, 0.7, 10, 0.6, "System Architecture & Data Flow", 32, WHITE, True)

# RAW layer
add_colored_card(slide, 0.8, 1.5, 11.7, 1.1, ORANGE)
add_text_box(slide, 1.1, 1.55, 4, 0.3, "RAW Schema - Data Ingestion", 16, ORANGE, True)
add_text_box(slide, 1.1, 1.9, 5, 0.5, "Landing zone. Daily task (2AM UTC) generates synthetic transactions. Stream captures CDC.", 12, GRAY)
add_text_box(slide, 7, 1.7, 5, 0.5, "RAW_CUSTOMERS  |  RAW_TRANSACTIONS  |  RAW_RISK_SCORES  |  RAW_ALERTS  |  STREAM", 11, ORANGE)

add_text_box(slide, 5, 2.7, 3, 0.3, "6 Dynamic Tables (auto-refresh, 1-min lag)", 12, BLUE, alignment=PP_ALIGN.CENTER)

# TRANSFORM layer
add_colored_card(slide, 0.8, 3.1, 11.7, 1.1, BLUE)
add_text_box(slide, 1.1, 3.15, 4, 0.3, "TRANSFORM Schema - Business Logic", 16, BLUE, True)
add_text_box(slide, 1.1, 3.5, 5, 0.5, "Enrichment: amount buckets, risk flags, SLA tracking, KYC actions, pattern detection.", 12, GRAY)
add_text_box(slide, 7, 3.3, 5, 0.5, "DT_CUSTOMERS  |  DT_TRANSACTIONS  |  DT_RISK_ENRICHED\nDT_ALERT_ENRICHED  |  DT_SUSPICIOUS_PATTERNS", 11, BLUE)

add_text_box(slide, 5, 4.3, 3, 0.3, "Semantic View + Cortex Search + AI Functions", 12, PURPLE, alignment=PP_ALIGN.CENTER)

# SEMANTIC layer
add_colored_card(slide, 0.8, 4.7, 11.7, 1.1, PURPLE)
add_text_box(slide, 1.1, 4.75, 4, 0.3, "SEMANTIC Schema - Intelligence Layer", 16, PURPLE, True)
add_text_box(slide, 1.1, 5.1, 5, 0.5, "Semantic view (16 dims, 4 metrics, 4 VQRs) + RAG (18 chunks) + Cortex Search.", 12, GRAY)
add_text_box(slide, 7, 4.9, 5, 0.5, "SV_RISK_INTELLIGENCE  |  SEARCH_SERVICE  |  AI_EXTRACT  |  AI_CLASSIFY", 11, PURPLE)

add_text_box(slide, 5, 5.9, 3, 0.3, "3 Cortex Agents + MCP Server + 5 Surfaces", 12, GREEN, alignment=PP_ALIGN.CENTER)

# APP layer
add_colored_card(slide, 0.8, 6.2, 11.7, 0.8, GREEN)
add_text_box(slide, 1.1, 6.25, 4, 0.3, "Application Layer - Multi-Surface", 16, GREEN, True)
add_text_box(slide, 7, 6.3, 5, 0.4, "STREAMLIT  |  CORTEX AGENTS (3)  |  MCP SERVER  |  COCO SKILLS (3)  |  SLACK BOT", 11, GREEN)

# ═══════════════════════════════════════════════
# SLIDE 4: COCO SKILLS + MULTI-AGENT
# ═══════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_text_box(slide, 0.8, 0.3, 4, 0.3, "COCO SKILLS & AGENTS", 11, PURPLE, True)
add_text_box(slide, 12, 0.3, 1, 0.3, "04 / 08", 10, DARK_GRAY, alignment=PP_ALIGN.RIGHT)
add_text_box(slide, 0.8, 0.7, 10, 0.6, "CoCo Skills & Multi-Agent Orchestration", 32, WHITE, True)

# Left - Skills
add_text_box(slide, 0.8, 1.5, 5, 0.3, "Reusable CoCo Skills", 18, GREEN, True)

add_colored_card(slide, 0.8, 2.0, 5.8, 1.2, GREEN)
add_text_box(slide, 1.1, 2.1, 5.3, 0.3, "/risk-copilot-deploy", 15, WHITE, True)
add_text_box(slide, 1.1, 2.5, 5.3, 0.5, "Deploys entire copilot to any Snowflake account: schemas, data, pipelines, agents, Streamlit app. Single command.", 12, GRAY)

add_colored_card(slide, 0.8, 3.4, 5.8, 1.2, BLUE)
add_text_box(slide, 1.1, 3.5, 5.3, 0.3, "/risk-copilot-investigate", 15, WHITE, True)
add_text_box(slide, 1.1, 3.9, 5.3, 0.5, "Investigates an alert or customer with RAG-backed AI evidence generation and audit logging.", 12, GRAY)

add_colored_card(slide, 0.8, 4.8, 5.8, 1.2, PURPLE)
add_text_box(slide, 1.1, 4.9, 5.3, 0.3, "/risk-copilot-orchestrate", 15, WHITE, True)
add_text_box(slide, 1.1, 5.3, 5.3, 0.5, "Coordinates 3 agents with clear handoffs: Triage -> Compliance -> Intelligence synthesis.", 12, GRAY)

# Right - Multi-Agent
add_text_box(slide, 7.2, 1.5, 5, 0.3, "Multi-Agent Architecture", 18, PURPLE, True)

add_colored_card(slide, 7.2, 2.0, 5.5, 1.1, BLUE)
add_text_box(slide, 7.5, 2.1, 5, 0.25, "RISK_TRIAGE_AGENT", 14, BLUE, True)
add_text_box(slide, 7.5, 2.4, 5, 0.5, "Data Specialist - cortex_analyst_text_to_sql\nScreens alerts, pulls transactions, risk scores", 12, GRAY)

add_text_box(slide, 9.5, 3.15, 2, 0.3, "handoff: triage context", 10, BLUE, alignment=PP_ALIGN.CENTER)

add_colored_card(slide, 7.2, 3.5, 5.5, 1.1, PURPLE)
add_text_box(slide, 7.5, 3.6, 5, 0.25, "REGULATORY_COMPLIANCE_AGENT", 14, PURPLE, True)
add_text_box(slide, 7.5, 3.9, 5, 0.5, "Policy Expert - cortex_search (RAG)\nFinds applicable regulations, filing requirements", 12, GRAY)

add_text_box(slide, 9.5, 4.65, 2, 0.3, "handoff: regulatory citations", 10, PURPLE, alignment=PP_ALIGN.CENTER)

add_colored_card(slide, 7.2, 5.0, 5.5, 1.1, GREEN)
add_text_box(slide, 7.5, 5.1, 5, 0.25, "RISK_INTELLIGENCE_AGENT", 14, GREEN, True)
add_text_box(slide, 7.5, 5.4, 5, 0.5, "Orchestrator - both tools combined\nSynthesizes data + regulations into evidence report", 12, GRAY)

# ═══════════════════════════════════════════════
# SLIDE 5: SOLUTION WALKTHROUGH
# ═══════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_text_box(slide, 0.8, 0.3, 4, 0.3, "SOLUTION WALKTHROUGH", 11, GREEN, True)
add_text_box(slide, 12, 0.3, 1, 0.3, "05 / 08", 10, DARK_GRAY, alignment=PP_ALIGN.RIGHT)
add_text_box(slide, 0.8, 0.7, 10, 0.6, "Signal -> Evidence -> Documented Finding", 32, WHITE, True)
add_text_box(slide, 0.8, 1.3, 10, 0.4, "Complete flow in a 5-page Streamlit copilot", 16, GRAY)

pages = [
    ("Dashboard", RED, "Color-coded KPI cards, severity badges, detection coverage, transaction volume trends, recent AI activity feed"),
    ("Investigation", BLUE, "Alert drill-down, customer profile cards, risk score gauge, transaction timeline, one-click evidence generation"),
    ("Regulatory Chat", PURPLE, "RAG-backed Q&A, styled chat bubbles, source cards with citations, confidence scoring (HIGH/MED/LOW)"),
    ("Audit Trail", GREEN, "Full log of every AI action: query, RAG context, LLM response, user, timestamp. Charts + export"),
    ("System Health", ORANGE, "Pipeline status, dynamic table freshness, task history, search service, agent/MCP monitoring"),
]

for i, (name, color, desc) in enumerate(pages):
    x = 0.8 + i * 2.45
    add_colored_card(slide, x, 1.9, 2.2, 3.2, color)
    add_text_box(slide, x + 0.15, 2.0, 1.9, 0.3, name, 15, color, True)
    add_text_box(slide, x + 0.15, 2.5, 1.9, 2.4, desc, 12, GRAY)

# Fraud patterns
add_colored_card(slide, 0.8, 5.4, 11.7, 1.5, CARD_BORDER)
add_text_box(slide, 1.1, 5.5, 10, 0.3, "9 Fraud & Risk Patterns Detected", 14, GRAY, True)
patterns = [
    "Structuring (sub-$10k deposits)", "Layering (multi-jurisdiction)", "Round-Tripping (circular flows)",
    "Velocity Anomaly (burst activity)", "Dormant Reactivation", "Shell Company Activity",
    "KYC Expired Activity", "PEP Large Transactions", "Mule Account Indicators"
]
for i, p in enumerate(patterns):
    col = i % 3
    row = i // 3
    add_text_box(slide, 1.1 + col * 4, 5.9 + row * 0.35, 3.5, 0.3, "  " + p, 11, WHITE)

# ═══════════════════════════════════════════════
# SLIDE 6: GUARDRAILS
# ═══════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_text_box(slide, 0.8, 0.3, 4, 0.3, "GUARDRAILS & GOVERNANCE", 11, GREEN, True)
add_text_box(slide, 12, 0.3, 1, 0.3, "06 / 08", 10, DARK_GRAY, alignment=PP_ALIGN.RIGHT)
add_text_box(slide, 0.8, 0.7, 10, 0.6, "Governed, Trustworthy AI Behavior", 32, WHITE, True)

# Left - AI Guardrails
add_text_box(slide, 0.8, 1.5, 5, 0.3, "AI Safety Guardrails", 18, GREEN, True)
guardrails = [
    ("Input Validation", "Blocks SQL injection (DROP, DELETE, TRUNCATE, MERGE), enforces 1000-char limit"),
    ("Topic Relevance", "Restricts to 20+ allowed topics (risk, fraud, AML, compliance, etc.)"),
    ("Confidence Scoring", "HIGH/MEDIUM/LOW based on regulatory citation count in LLM output"),
    ("RAG Grounding", "LLM uses ONLY regulatory context. Flags uncertainty explicitly."),
    ("Output Validation", "Detects low-confidence phrases, missing citations, short responses"),
    ("Blocked Keywords", "11 dangerous SQL patterns blocked at input layer (read-only copilot)"),
    ("Error Handling", "Cortex API failures return user-friendly messages, never raw traces"),
]
for i, (title, desc) in enumerate(guardrails):
    y = 2.0 + i * 0.7
    add_text_box(slide, 1.1, y, 5, 0.25, title, 14, WHITE, True)
    add_text_box(slide, 1.1, y + 0.28, 5, 0.4, desc, 11, GRAY)

# Right - Data Governance
add_text_box(slide, 7.2, 1.5, 5, 0.3, "Data Governance", 18, PURPLE, True)

add_colored_card(slide, 7.2, 2.0, 5.5, 1.1, BLUE)
add_text_box(slide, 7.5, 2.1, 5, 0.25, "RBAC - Role-Based Access", 14, WHITE, True)
add_text_box(slide, 7.5, 2.4, 5, 0.5, "RISK_ANALYST: Investigate + generate evidence\nRISK_AUDITOR: Read-only access + audit trail", 12, GRAY)

add_colored_card(slide, 7.2, 3.3, 5.5, 1.3, PURPLE)
add_text_box(slide, 7.5, 3.4, 5, 0.25, "PII Protection", 14, WHITE, True)
add_text_box(slide, 7.5, 3.7, 5, 0.7, "Dynamic Masking: PII_MASK policy on name, email, phone\nSemantic Tags: SNOWFLAKE.CORE.SEMANTIC_CATEGORY\nRISK_AUDITOR sees: ***MASKED***", 12, GRAY)

add_colored_card(slide, 7.2, 4.8, 5.5, 1.0, GREEN)
add_text_box(slide, 7.5, 4.9, 5, 0.25, "Complete Audit Trail", 14, WHITE, True)
add_text_box(slide, 7.5, 5.2, 5, 0.5, "Every AI interaction logged: action type, alert ID, user query, RAG context, LLM response, user, timestamp", 12, GRAY)

add_colored_card(slide, 7.2, 6.0, 5.5, 0.9, ORANGE)
add_text_box(slide, 7.5, 6.1, 5, 0.25, "Human-in-the-Loop Policy", 14, WHITE, True)
add_text_box(slide, 7.5, 6.4, 5, 0.4, "AI outputs are advisory only. Human investigator must review and approve before SAR filing.", 12, GRAY)

# ═══════════════════════════════════════════════
# SLIDE 7: IMPACT STATEMENT
# ═══════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_text_box(slide, 0.8, 0.3, 4, 0.3, "IMPACT STATEMENT", 11, BLUE, True)
add_text_box(slide, 12, 0.3, 1, 0.3, "07 / 08", 10, DARK_GRAY, alignment=PP_ALIGN.RIGHT)
add_text_box(slide, 0.8, 0.7, 10, 0.6, "Measurable Impact & Scalability", 32, WHITE, True)

# KPI cards
kpis = [
    ("85%", "Faster Investigation", "4-8 hours -> 30-45 min", BLUE),
    ("100%", "Audit Coverage", "Every AI interaction logged", GREEN),
    ("5x", "More Alerts Processed", "Same team, 5x throughput", PURPLE),
    ("0", "Data Movement", "All native in Snowflake", ORANGE),
]
for i, (num, label, desc, color) in enumerate(kpis):
    x = 0.8 + i * 3.1
    add_colored_card(slide, x, 1.5, 2.8, 1.5, color)
    add_text_box(slide, x + 0.2, 1.6, 2.4, 0.6, num, 36, color, True, PP_ALIGN.CENTER)
    add_text_box(slide, x + 0.2, 2.2, 2.4, 0.3, label, 13, WHITE, True, PP_ALIGN.CENTER)
    add_text_box(slide, x + 0.2, 2.5, 2.4, 0.3, desc, 11, GRAY, alignment=PP_ALIGN.CENTER)

# Before/After
add_text_box(slide, 0.8, 3.3, 5, 0.3, "BEFORE vs AFTER", 11, GRAY, True)
befores = [
    "Manual regulatory document search (30+ min)",
    "Copy-paste evidence from 5+ systems into Word doc",
    "No audit trail for AI-assisted decisions",
    "Siloed tools: email, spreadsheets, document folders",
]
afters = [
    "Instant RAG-powered policy lookup with citations",
    "One-click structured evidence report",
    "Full explainability: query + context + response logged",
    "5 unified surfaces: Streamlit, Agent, CoCo, MCP, Slack",
]
for i in range(4):
    y = 3.7 + i * 0.65
    add_text_box(slide, 0.8, y, 5, 0.3, befores[i], 12, RED)
    add_text_box(slide, 6.2, y, 0.5, 0.3, "->", 14, BLUE, True, PP_ALIGN.CENTER)
    add_text_box(slide, 7, y, 5.5, 0.3, afters[i], 12, GREEN, True)

# Scalability
add_text_box(slide, 0.8, 6.3, 10, 0.3, "SCALABILITY & EXTENSIBILITY", 11, GRAY, True)
scale_items = [
    ("Production Scale", "Dynamic tables handle millions of txns"),
    ("New Regulations", "AI_EXTRACT + AI_CLASSIFY auto-process new docs"),
    ("Single-Script Deploy", "deploy_all.sql - paste and run on any account"),
]
for i, (t, d) in enumerate(scale_items):
    x = 0.8 + i * 4.1
    add_colored_card(slide, x, 6.6, 3.8, 0.7, CARD_BORDER)
    add_text_box(slide, x + 0.15, 6.65, 3.5, 0.2, t, 12, WHITE, True)
    add_text_box(slide, x + 0.15, 6.9, 3.5, 0.3, d, 11, GRAY)

# ═══════════════════════════════════════════════
# SLIDE 8: SUMMARY
# ═══════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_text_box(slide, 0.8, 0.3, 4, 0.3, "COCO USAGE & SUMMARY", 11, GREEN, True)
add_text_box(slide, 12, 0.3, 1, 0.3, "08 / 08", 10, DARK_GRAY, alignment=PP_ALIGN.RIGHT)
add_text_box(slide, 0.8, 0.7, 10, 0.6, "Built 100% with CoCo Desktop", 32, WHITE, True)

# CoCo phases
phases = [
    ("Planning", "Data model, ontology, workflow design"),
    ("Development", "All SQL, Python, YAML authored in CoCo"),
    ("Execution", "SQL run, app deployed, tasks scheduled"),
    ("Testing", "8 validation tests + guardrail checks"),
]
for i, (name, desc) in enumerate(phases):
    x = 0.8 + i * 3.1
    add_colored_card(slide, x, 1.5, 2.8, 1.0, CARD_BORDER)
    add_text_box(slide, x + 0.2, 1.55, 2.4, 0.3, name, 15, WHITE, True, PP_ALIGN.CENTER)
    add_text_box(slide, x + 0.2, 1.9, 2.4, 0.4, desc, 11, GRAY, alignment=PP_ALIGN.CENTER)

# Checklist
add_colored_card(slide, 0.8, 2.8, 11.7, 1.5, CARD_BORDER)
add_text_box(slide, 1.1, 2.9, 10, 0.3, "Hackathon Criteria - All Complete", 14, WHITE, True, PP_ALIGN.CENTER)
checklist = [
    "Synthetic Data (20 customers, 50 txns)", "Data Pipeline (6 DTs, 3 tasks)", "Semantic View (4 VQRs)", "Streamlit App (5 pages)",
    "MCP Server", "Document Processing", "Reusable Skills (3)", "Automations (3 scheduled tasks)",
    "Multi-Agent (3 agents)", "Multi-Surface (5)", "AI Guardrails (7 checks)", "Slack Bot"
]
for i, item in enumerate(checklist):
    col = i % 4
    row = i // 4
    x = 1.5 + col * 2.8
    y = 3.35 + row * 0.35
    add_text_box(slide, x, y, 2.5, 0.3, "  " + item, 12, GREEN, True)

# Stats
stats = [("40+", "Snowflake Objects", BLUE), ("3", "Cortex Agents", PURPLE), ("5", "Access Surfaces", GREEN), ("7", "Guardrail Checks", ORANGE)]
for i, (num, label, color) in enumerate(stats):
    x = 0.8 + i * 3.1
    add_colored_card(slide, x, 4.6, 2.8, 1.0, color)
    add_text_box(slide, x + 0.2, 4.65, 2.4, 0.5, num, 32, color, True, PP_ALIGN.CENTER)
    add_text_box(slide, x + 0.2, 5.15, 2.4, 0.3, label, 12, GRAY, alignment=PP_ALIGN.CENTER)

# Thank you
add_text_box(slide, 0.8, 5.9, 11.7, 0.5, "Thank You", 24, WHITE, True, PP_ALIGN.CENTER)
add_text_box(slide, 0.8, 6.4, 11.7, 0.3, "GitHub: github.com/rdevath21/risk-fraud-regulatory-intelligence-copilot", 14, GRAY, alignment=PP_ALIGN.CENTER)
add_text_box(slide, 0.8, 6.8, 11.7, 0.3, "Single-script deployment: deploy/deploy_all.sql  |  Verified on 2 Snowflake accounts  |  Streamlit: RISK_FRAUD_COPILOT", 12, DARK_GRAY, alignment=PP_ALIGN.CENTER)

# Save
output_path = r"c:\Users\rdeva\.snowflake\cortex\playground\workspace\presentation\Risk_Copilot_Submission_Deck_v2.pptx"
prs.save(output_path)
print(f"PPTX saved to: {output_path}")
